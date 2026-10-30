from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
source = ROOT / 'local-data/portfolio.json'
if not source.exists(): source = ROOT / 'data/portfolio.seed.json'
data = json.loads(source.read_text(encoding='utf-8'))
projects = [p for p in data['projects'] if p.get('visible')]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.76), Inches(.7)
sec.left_margin = sec.right_margin = Inches(.82)
sec.header_distance = sec.footer_distance = Inches(.32)
for key in ['Normal','Title','Heading 1','Heading 2','Header','Footer']:
    s = doc.styles[key]
    s.font.name = 'Microsoft YaHei'
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    s.font.color.rgb = RGBColor(0,0,0)
normal=doc.styles['Normal']
normal.font.size=Pt(13)
normal.paragraph_format.line_spacing=1.55
normal.paragraph_format.space_after=Pt(10)
doc.styles['Title'].font.size=Pt(25)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(14)
header=sec.header.paragraphs[0]
header.text='KUZ  /  工程项目介绍'
header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Kuz    |    ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
footer.add_run(' / '+str(len(projects))).font.size=Pt(9)

def para(text,size=13,bold=False,after=10):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(after)
    r=p.add_run(text);r.font.size=Pt(size);r.bold=bold
    return p
def label(text):
    p=para(text,12,True,5);p.paragraph_format.space_before=Pt(18);p.paragraph_format.keep_with_next=True
def link(url):
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(16)
    p.add_run('GitHub  ').font.size=Pt(9)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz)
    color=OxmlElement('w:color');color.set(qn('w:val'),'334B45');pr.append(color)
    r.append(pr);t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)

for i,p in enumerate(projects):
    if i: doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    title=doc.add_paragraph(f'工程项目 {i+1:02d}',style='Title')
    if i==0:
        para(f'Kuz 的工程项目介绍。覆盖当前作品集的 {len(projects)} 个仓库条目，逐项说明技术、功能与所针对的问题；学习与原型项目保留实际状态。',10,False,8)
    label('工程名称：')
    para(p['name']['zh'],18,True,7)
    label('工程技术栈：')
    para(' / '.join(p['tech']) or '无可核验的软件技术栈',12.5)
    label('工程简单描述：')
    para(p['summary']['zh'])
    label('解决的痛点：')
    para(p['problem']['zh'],13,False,8)
    para(p['solution']['zh'])
    para('项目状态：'+p['status']['zh'],10,False,3)
    link(p['github'])

doc.core_properties.title='Kuz 工程项目介绍'
doc.core_properties.author='Kuz'
doc.core_properties.subject=f'{len(projects)} 个工程的名称 技术栈 简介与解决的痛点'
for tree in (doc.styles.element,doc.element):
    for border in list(tree.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
output=ROOT/'output/docx/Kuz_工程项目介绍.docx'
output.parent.mkdir(parents=True,exist_ok=True)
doc.save(output)
print(json.dumps({'path':str(output),'projects':len(projects),'explicit_page_breaks':len(projects)-1},ensure_ascii=False))
